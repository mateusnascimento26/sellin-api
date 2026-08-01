import io

import openpyxl
from openpyxl.utils.exceptions import InvalidFileException

from domain.exceptions import SpreadsheetReadError

EXPECTED_SHEET_NAME = "Export"


class OpenPyXLSpreadsheetReader:
    def read_export_sheet(self, content: bytes) -> tuple[list[str], list[list[object]]]:
        try:
            workbook = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
        except InvalidFileException:
            raise SpreadsheetReadError("O arquivo não é um arquivo Excel (.xlsx) válido.")
        except Exception:
            raise SpreadsheetReadError("O arquivo está corrompido ou protegido por senha e não pôde ser aberto.")

        if EXPECTED_SHEET_NAME not in workbook.sheetnames:
            raise SpreadsheetReadError(f"A aba '{EXPECTED_SHEET_NAME}' não foi encontrada no arquivo.")

        sheet = workbook[EXPECTED_SHEET_NAME]
        rows_iterator = sheet.iter_rows(values_only=True)

        header_row = next(rows_iterator, None)
        if header_row is None:
            raise SpreadsheetReadError(f"A aba '{EXPECTED_SHEET_NAME}' está vazia.")

        data_rows = [list(row) for row in rows_iterator]
        return list(header_row), data_rows